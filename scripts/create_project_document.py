from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "BikeShop_Project_Overview.docx"
BLUE = "1F4E78"
PALE = "EAF2F8"
GRID = "D9D9D9"


def set_cell_fill(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    repeat_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_fill(cell, BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(9)
        if widths:
            cell.width = Cm(widths[i])
        set_cell_margins(cell)
    for r_index, values in enumerate(rows):
        row = table.add_row()
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.text = str(value)
            if r_index % 2 == 1:
                set_cell_fill(cell, PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.5)
            if widths:
                cell.width = Cm(widths[i])
            set_cell_margins(cell)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.add_run(item)


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.8)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.0)
section.right_margin = Cm(2.0)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(6)
styles["Normal"].paragraph_format.line_spacing = 1.08
for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12)):
    style = styles[style_name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(7)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(5)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run("Tài liệu tổng quan dự án BikeShop SWP391")
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run("Phạm vi, chức năng, kiến trúc và kế hoạch triển khai")
run.bold = True
run.font.size = Pt(13)
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Nhóm 4  |  Lead Hà Thanh Huy  |  Công nghệ .NET 10 và SQL Server").italic = True

doc.add_paragraph(
    "Tài liệu này là baseline ban đầu để nhóm thống nhất trước khi viết SRS và SDS. "
    "Dự án xây dựng website bán xe đạp và phụ kiện ở quy mô môn học, tập trung vào ba luồng có thể thiết kế, lập trình, kiểm thử và bảo vệ rõ ràng trong mười tuần."
)

doc.add_heading("1 Mục tiêu và phạm vi", level=1)
doc.add_paragraph(
    "BikeShop tham khảo cách tổ chức catalog của xedap.vn nhưng sử dụng nội dung và giao diện riêng. "
    "Hệ thống phục vụ Guest, Customer, Staff và Admin; thanh toán COD là phương thức duy nhất trong phạm vi chính. "
    "Guest là người chưa đăng nhập, không phải role tài khoản trong database."
)
add_table(doc, ["Trong phạm vi", "Ngoài phạm vi ban đầu"], [[
    "Catalog, biến thể, tìm kiếm, lọc, giỏ hàng, checkout COD, quản lý đơn, tồn kho và dashboard",
    "Thanh toán online, voucher phức tạp, chat, đánh giá, OAuth, vận chuyển bên thứ ba, đa chi nhánh và AI recommendation"
]], [8.2, 8.2])

doc.add_heading("2 Actor và workflow", level=1)
add_table(doc, ["Actor", "Quyền chính"], [
    ["Guest", "Xem trang chủ, danh sách, chi tiết, tìm kiếm và lọc sản phẩm"],
    ["Customer", "Quản lý tài khoản, địa chỉ, giỏ hàng, đặt hàng COD, xem và hủy đơn hợp lệ"],
    ["Staff", "Xử lý đơn hàng, cập nhật trạng thái, điều chỉnh tồn kho và xem cảnh báo sắp hết hàng"],
    ["Admin", "Quản lý catalog, tài khoản nhân viên, phân quyền, toàn bộ đơn hàng và dashboard"],
], [3.4, 13.0])

doc.add_heading("2.1 Workflow 0 Khởi tạo dữ liệu", level=2)
doc.add_paragraph("Admin tạo danh mục và thương hiệu, sau đó tạo sản phẩm, biến thể, hình ảnh và tồn kho ban đầu. Mỗi điều chỉnh tồn kho tạo một bản ghi lịch sử.")
doc.add_heading("2.2 Workflow 1 Mua hàng", level=2)
doc.add_paragraph("Xem hoặc tìm sản phẩm → chọn biến thể → thêm giỏ → nhập địa chỉ → xác nhận COD → tạo đơn → trừ tồn kho.")
add_bullets(doc, [
    "Từ chối checkout khi biến thể hết hàng hoặc số lượng vượt tồn kho.",
    "Kiểm tra lại trạng thái và giá của sản phẩm trước khi tạo đơn.",
    "Tạo Order, OrderItem, InventoryTransaction và OrderStatusHistory trong cùng transaction.",
    "Khi đơn được hủy hợp lệ, hoàn tồn kho và ghi lịch sử."
])
doc.add_heading("2.3 Workflow 2 Xử lý đơn và dashboard", level=2)
doc.add_paragraph("Staff xử lý Pending → Confirmed → Packing → Shipping → Completed. Đơn Pending hoặc Confirmed có thể chuyển sang Cancelled theo quyền và business rule. Admin có thể giám sát toàn bộ hàng đợi và dashboard.")

doc.add_page_break()
doc.add_heading("3 Danh sách chức năng và phân công", level=1)
function_rows = [
    ["Nguyễn Phước Hậu", "Account", "F01", "Đăng ký tài khoản và validation"],
    ["Nguyễn Phước Hậu", "Account", "F02", "Đăng nhập và đăng xuất"],
    ["Nguyễn Phước Hậu", "Account", "F03", "Xem và cập nhật hồ sơ"],
    ["Nguyễn Phước Hậu", "Account", "F04", "Đổi mật khẩu"],
    ["Nguyễn Phước Hậu", "Account", "F05", "Quản lý địa chỉ giao hàng"],
    ["Nguyễn Phước Hậu", "Account", "F06", "Admin tạo, khóa, mở và gán role Staff cho tài khoản nội bộ"],
    ["Hà Thanh Huy", "Catalog", "F07", "Trang chủ và nhóm sản phẩm nổi bật"],
    ["Hà Thanh Huy", "Catalog", "F08", "Danh sách sản phẩm và phân trang"],
    ["Hà Thanh Huy", "Catalog", "F09", "Chi tiết và chọn biến thể"],
    ["Hà Thanh Huy", "Catalog", "F10", "Tìm kiếm theo từ khóa"],
    ["Hà Thanh Huy", "Catalog", "F11", "Lọc theo danh mục, thương hiệu và giá"],
    ["Hà Thanh Huy", "Catalog", "F12", "Sắp xếp kết quả"],
    ["Trần Vũ Khang", "Cart Checkout", "F13", "Thêm biến thể vào giỏ"],
    ["Trần Vũ Khang", "Cart Checkout", "F14", "Xem giỏ hàng"],
    ["Trần Vũ Khang", "Cart Checkout", "F15", "Cập nhật số lượng"],
    ["Trần Vũ Khang", "Cart Checkout", "F16", "Xóa sản phẩm khỏi giỏ"],
    ["Trần Vũ Khang", "Cart Checkout", "F17", "Kiểm tra giá và tồn kho"],
    ["Trần Vũ Khang", "Cart Checkout", "F18", "Checkout COD và tạo đơn"],
    ["Huỳnh Vương Khánh", "Admin Catalog", "F19", "Quản lý danh mục"],
    ["Huỳnh Vương Khánh", "Admin Catalog", "F20", "Quản lý thương hiệu"],
    ["Huỳnh Vương Khánh", "Admin Catalog", "F21", "Quản lý sản phẩm"],
    ["Huỳnh Vương Khánh", "Admin Catalog", "F22", "Quản lý biến thể"],
    ["Huỳnh Vương Khánh", "Admin Catalog", "F23", "Quản lý hình ảnh"],
    ["Huỳnh Vương Khánh", "Catalog Inventory", "F24", "Staff hoặc Admin điều chỉnh và xem lịch sử tồn kho"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F25", "Khách xem lịch sử và chi tiết đơn"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F26", "Khách hủy đơn hợp lệ"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F27", "Staff hoặc Admin tìm và lọc hàng đợi đơn"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F28", "Staff chuyển trạng thái và hệ thống ghi lịch sử"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F29", "Dashboard doanh thu và đơn hàng"],
    ["Phạm Đình Đăng Khoa", "Orders Reports", "F30", "Staff hoặc Admin xem sản phẩm sắp hết hàng"],
]
add_table(doc, ["Owner", "Module", "ID", "Chức năng"], function_rows, [4.0, 3.0, 1.2, 8.2])
doc.add_paragraph("Danh sách F01–F30 là bản đề xuất để giảng viên xác nhận cách tính function. CRUD của cùng một entity thuộc cùng một owner.")

doc.add_heading("4 Công nghệ và kiến trúc", level=1)
add_table(doc, ["Thành phần", "Quyết định"], [
    ["Web", "ASP.NET Core MVC, Razor Views và Razor Pages của Identity"],
    ["Runtime", ".NET SDK 10.0.401, khóa bằng global.json"],
    ["Data", "Entity Framework Core 10 và SQL Server 2022 Express"],
    ["Security", "ASP.NET Core Identity, role Customer, Staff và Admin; Guest là người chưa xác thực"],
    ["Frontend", "Bootstrap 5 và JavaScript thuần"],
    ["Testing", "xUnit, unit test business rules và integration test với SQL Server test database"],
    ["Collaboration", "GitHub issues, feature branches, pull request review và milestone tags"],
], [4.0, 12.4])
doc.add_paragraph("Luồng phụ thuộc chuẩn: Web → Application → Domain; Infrastructure triển khai data access và được Web đăng ký bằng dependency injection. Không đưa Entity Framework Core vào Domain.")

doc.add_heading("5 Thiết kế database và data dictionary", level=1)
doc.add_paragraph(
    "Tất cả khóa chính kiểu int hoặc bigint của bảng nghiệp vụ được SQL Server sinh tự động bằng IDENTITY(1,1). "
    "Riêng AspNetUsers.Id và AspNetRoles.Id là nvarchar(450) vì ASP.NET Core Identity dùng khóa chuỗi; ứng dụng không yêu cầu người dùng tự nhập các giá trị này."
)
db_rows = [
    ["Categories", "Id int IDENTITY PK\nName nvarchar(100)\nSlug varchar(120) UQ\nDescription nvarchar(500) NULL\nIsActive bit", "Danh mục sản phẩm"],
    ["Brands", "Id int IDENTITY PK\nName nvarchar(100)\nSlug varchar(120) UQ\nIsActive bit", "Thương hiệu"],
    ["Products", "Id int IDENTITY PK\nCategoryId int FK\nBrandId int FK\nName nvarchar(200)\nSlug varchar(220) UQ\nDescription nvarchar(max) NULL\nIsActive bit\nCreatedAt datetime2", "Thông tin chung; ngừng bán bằng IsActive"],
    ["ProductVariants", "Id int IDENTITY PK\nProductId int FK\nSku varchar(50) UQ\nColor nvarchar(50) NULL\nSize nvarchar(30) NULL\nPrice decimal(18,2)\nStockQuantity int\nIsActive bit\nRowVersion rowversion", "Đơn vị thực tế được bán và giữ tồn kho"],
    ["ProductImages", "Id int IDENTITY PK\nProductId int FK\nImageUrl nvarchar(500)\nAltText nvarchar(200) NULL\nSortOrder int", "Nhiều ảnh cho một sản phẩm"],
    ["InventoryTransactions", "Id bigint IDENTITY PK\nProductVariantId int FK\nQuantityChange int\nReason nvarchar(100)\nReferenceCode nvarchar(50) NULL\nCreatedByUserId nvarchar(450) FK NULL\nCreatedAt datetime2", "Audit nhập, trừ và hoàn kho"],
    ["Addresses", "Id int IDENTITY PK\nUserId nvarchar(450) FK\nRecipientName nvarchar(120)\nPhoneNumber varchar(20)\nAddressLine nvarchar(250)\nWard nvarchar(100)\nDistrict nvarchar(100)\nProvince nvarchar(100)\nIsDefault bit", "Địa chỉ lưu của khách hàng"],
    ["Carts", "Id int IDENTITY PK\nUserId nvarchar(450) FK UQ\nUpdatedAt datetime2", "Một giỏ hiện tại cho mỗi khách"],
    ["CartItems", "Id int IDENTITY PK\nCartId int FK\nProductVariantId int FK\nQuantity int", "UQ CartId + ProductVariantId; Quantity > 0"],
    ["Orders", "Id bigint IDENTITY PK\nOrderCode varchar(30) UQ\nUserId nvarchar(450) FK\nStatus int\nRecipientName nvarchar(120)\nPhoneNumber varchar(20)\nShippingAddress nvarchar(600)\nSubtotal decimal(18,2)\nShippingFee decimal(18,2)\nTotalAmount decimal(18,2)\nCustomerNote nvarchar(500) NULL\nCreatedAt, UpdatedAt datetime2", "Đơn hàng và snapshot địa chỉ giao"],
    ["OrderItems", "Id bigint IDENTITY PK\nOrderId bigint FK\nProductVariantId int FK\nProductName nvarchar(200)\nSku varchar(50)\nVariantDescription nvarchar(120) NULL\nUnitPrice decimal(18,2)\nQuantity int\nLineTotal decimal(18,2)", "Snapshot sản phẩm và giá lúc mua"],
    ["OrderStatusHistories", "Id bigint IDENTITY PK\nOrderId bigint FK\nOldStatus int\nNewStatus int\nNote nvarchar(500) NULL\nChangedByUserId nvarchar(450) FK NULL\nChangedAt datetime2", "Audit mọi lần đổi trạng thái"],
    ["AspNetUsers", "Id nvarchar(450) PK\nFullName nvarchar(150) NULL\nIsActive bit\nCreatedAt datetime2\nUserName, Email, PasswordHash và các cột Identity", "Identity sinh khóa chuỗi và hash mật khẩu"],
    ["AspNetRoles", "Id nvarchar(450) PK\nName nvarchar(256)\nNormalizedName nvarchar(256) UQ\nConcurrencyStamp nvarchar(max)", "Ba role tài khoản Customer, Staff và Admin"],
    ["Identity mapping", "AspNetUserRoles\nAspNetUserClaims\nAspNetRoleClaims\nAspNetUserLogins\nAspNetUserTokens", "Bảng liên kết và xác thực do Identity quản lý"],
]
add_table(doc, ["Bảng", "Cột và kiểu dữ liệu", "Mục đích và ràng buộc"], db_rows, [3.2, 8.0, 5.2])
add_bullets(doc, [
    "Giá dùng decimal(18,2); SKU, slug và mã đơn có unique constraint.",
    "Sản phẩm đã phát sinh đơn được ngừng bán bằng IsActive thay vì hard delete.",
    "OrderItem lưu snapshot để lịch sử đơn không thay đổi khi catalog được sửa.",
    "Checkout dùng transaction và kiểm tra RowVersion để hạn chế oversell.",
    "DATA_CREATE.sql tạo schema; DATA_INSERT.sql thêm role và catalog demo có thể chạy lại.",
    "Entity Framework Core được cấu hình ValueGeneratedOnAdd cho toàn bộ khóa số để khớp IDENTITY(1,1)."
])

doc.add_heading("5.1 Ma trận quyền", level=2)
add_table(doc, ["Nghiệp vụ", "Guest", "Customer", "Staff", "Admin"], [
    ["Xem catalog", "Có", "Có", "Có", "Có"],
    ["Giỏ hàng và đặt hàng", "Không", "Có", "Không", "Không"],
    ["Xem và hủy đơn của mình", "Không", "Có", "Không", "Không"],
    ["Xử lý đơn và cập nhật trạng thái", "Không", "Không", "Có", "Có"],
    ["Điều chỉnh tồn kho", "Không", "Không", "Có", "Có"],
    ["Quản lý catalog", "Không", "Không", "Không", "Có"],
    ["Quản lý tài khoản Staff và phân quyền", "Không", "Không", "Không", "Có"],
    ["Dashboard và báo cáo tổng hợp", "Không", "Không", "Không", "Có"],
], [5.8, 2.2, 2.6, 2.6, 2.6])

doc.add_heading("6 Yêu cầu chất lượng", level=1)
add_table(doc, ["Mục", "Tiêu chí chấp nhận"], [
    ["Security", "Authorization theo role; server-side validation; không commit secret"],
    ["Data integrity", "PK, FK, unique, check constraint; checkout atomic"],
    ["Usability", "Responsive ở desktop và mobile; thông báo lỗi rõ ràng"],
    ["Performance", "Catalog phân trang; truy vấn đọc dùng projection và AsNoTracking khi phù hợp"],
    ["Auditability", "Lưu lịch sử đơn, tồn kho; issue và pull request liên kết chức năng"],
    ["Testing", "Ít nhất 25 test case toàn nhóm; ưu tiên core workflow và exception path"],
], [4.0, 12.4])

doc.add_heading("7 Kế hoạch mười tuần", level=1)
plan_rows = [
    ["1", "Scope, actor, business rule, user story, Git và skeleton"],
    ["2", "Wireframe, ERD, architecture, state diagram và traceability"],
    ["3", "Assessment 1 và một vertical slice chạy được"],
    ["4", "Identity, seed data, category, brand và product"],
    ["5", "Catalog khách hàng, variant, image và inventory"],
    ["6", "Cart, checkout, transaction và exception path"],
    ["7", "Order processing, integration và migration freeze cho A2"],
    ["8", "SDS, sequence diagrams và Assessment 2"],
    ["9", "Dashboard, test, sửa lỗi và thử deployment"],
    ["10", "Assessment 3, regression test, evidence và rehearsal demo"],
]
add_table(doc, ["Tuần", "Đầu ra chính"], plan_rows, [2.0, 14.4])

doc.add_heading("8 Quy tắc Git và Definition of Done", level=1)
add_numbered(doc, [
    "Tạo issue có requirement và acceptance criteria trước khi code.",
    "Tạo nhánh feature/<issue-id>-<ten-ngan> từ develop; không push trực tiếp main.",
    "Code phải build, có validation, xử lý exception path và cập nhật database migration khi cần.",
    "Tác giả tự test; pull request có ít nhất một reviewer và liên kết issue.",
    "Cập nhật traceability từ function đến màn hình, class, database, test case và commit.",
    "Chỉ merge khi dotnet build và dotnet test đều thành công."
])

doc.add_heading("9 Baseline cần duyệt trong buổi họp đầu", level=1)
add_bullets(doc, [
    "Tên dự án và phạm vi ngoài scope.",
    "Danh sách F01–F30 và cách giảng viên tính function cá nhân.",
    "Business rule trừ kho khi tạo đơn và hoàn kho khi hủy.",
    "Chuẩn máy: .NET 10.0.401, SQL Server 2022 Express, SSMS 21 và Visual Studio stable.",
    "Phân công owner, người điều phối migration và lịch review hằng tuần."
])

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
